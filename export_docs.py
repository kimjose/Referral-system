#!/usr/bin/env python3

import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import markdown
import re
from bs4 import BeautifulSoup
import requests
from urllib.parse import urljoin, urlparse
import glob

def setup_document_styles(doc):
    # Add Code style
    styles = doc.styles
    try:
        code_style = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
        code_style.base_style = styles['Normal']
        code_font = code_style.font
        code_font.name = 'Courier New'
        code_font.size = Pt(9)
    except ValueError:
        # Style already exists
        pass

def find_image_file(image_name, base_dir):
    """Find an image file in various possible locations"""
    # Remove any query parameters or fragments from the image name
    image_name = image_name.split('?')[0].split('#')[0]
    
    # Try different possible paths
    possible_paths = [
        os.path.join(base_dir, image_name),  # Direct path
        os.path.join(base_dir, 'images', image_name),  # images subdirectory
        os.path.join(os.path.dirname(base_dir), 'images', image_name),  # Parent images directory
        os.path.join('docs', 'images', image_name),  # docs/images directory
        os.path.join('images', image_name),  # Root images directory
    ]
    
    # Also try with different extensions if the original doesn't exist
    if not any(os.path.exists(p) for p in possible_paths):
        base_name = os.path.splitext(image_name)[0]
        for ext in ['.png', '.jpg', '.jpeg', '.gif', '.svg']:
            possible_paths.extend([
                os.path.join(base_dir, base_name + ext),
                os.path.join(base_dir, 'images', base_name + ext),
                os.path.join(os.path.dirname(base_dir), 'images', base_name + ext),
                os.path.join('docs', 'images', base_name + ext),
                os.path.join('images', base_name + ext),
            ])
    
    # Try to find the image
    for path in possible_paths:
        if os.path.exists(path):
            return path
    
    # If not found, try to find any image with a similar name
    for path in possible_paths:
        dir_name = os.path.dirname(path)
        if os.path.exists(dir_name):
            base_name = os.path.splitext(os.path.basename(path))[0]
            similar_files = glob.glob(os.path.join(dir_name, f"{base_name}*"))
            if similar_files:
                return similar_files[0]
    
    return None

def process_image(src, md_dir, doc):
    """Process an image source and add it to the document"""
    print(f"Processing image: {src}")
    
    # Handle remote images
    if src.startswith(('http://', 'https://')):
        try:
            response = requests.get(src)
            if response.status_code == 200:
                img_path = os.path.join(md_dir, 'temp_image.jpg')
                with open(img_path, 'wb') as f:
                    f.write(response.content)
                doc.add_picture(img_path, width=Inches(6))
                os.remove(img_path)  # Clean up temporary file
                print(f"Successfully added remote image: {src}")
            else:
                print(f"Failed to download remote image: {src} (Status code: {response.status_code})")
        except Exception as e:
            print(f"Error downloading remote image {src}: {e}")
        return

    # Handle local images
    img_path = find_image_file(src, md_dir)
    if img_path:
        try:
            doc.add_picture(img_path, width=Inches(6))
            print(f"Successfully added local image: {img_path}")
        except Exception as e:
            print(f"Error adding local image {img_path}: {e}")
    else:
        print(f"Image not found: {src}")

def convert_md_to_docx(md_file, doc):
    """Convert a markdown file to Word document sections"""
    print(f"\nProcessing file: {md_file}")
    
    with open(md_file, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Convert markdown to HTML
    html = markdown.markdown(md_content, extensions=['tables', 'fenced_code'])
    
    # Parse HTML
    soup = BeautifulSoup(html, 'html.parser')
    
    # Get the directory of the markdown file for resolving relative image paths
    md_dir = os.path.dirname(os.path.abspath(md_file))
    print(f"Markdown file directory: {md_dir}")
    
    # Process each element
    for element in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'pre', 'img']):
        if element.name.startswith('h'):
            level = int(element.name[1])
            doc.add_heading(element.get_text(), level=level)
        elif element.name == 'p':
            # Check if paragraph contains an image
            img = element.find('img')
            if img:
                src = img.get('src')
                if src:
                    process_image(src, md_dir, doc)
            else:
                doc.add_paragraph(element.get_text())
        elif element.name == 'pre':
            # Handle code blocks
            code = element.get_text()
            p = doc.add_paragraph()
            p.style = 'Code'
            p.text = code
        elif element.name == 'img':
            src = element.get('src')
            if src:
                process_image(src, md_dir, doc)

def create_documentation():
    # Create a new Word document
    doc = Document()
    
    # Setup document styles
    setup_document_styles(doc)
    
    # Add title
    title = doc.add_heading('Angaza Referral System Documentation', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add introduction
    doc.add_paragraph('Welcome to the Angaza Referral System documentation. This documentation provides comprehensive information about the system\'s features, setup, and usage.')
    
    # Process each section
    sections = {
        'User Guide': 'user-guide',
        'Development': 'development',
        'Architecture': 'architecture',
        'API Reference': 'api',
        'Standards': 'standards'
    }
    
    for section_name, section_dir in sections.items():
        # Add section heading
        doc.add_heading(section_name, level=1)
        
        # Process each file in the section
        section_path = os.path.join('docs', section_dir)
        if os.path.exists(section_path):
            for file in os.listdir(section_path):
                if file.endswith('.md'):
                    file_path = os.path.join(section_path, file)
                    convert_md_to_docx(file_path, doc)
    
    # Add support section
    doc.add_heading('Support', level=1)
    support_text = """
For additional support:
- Email: support@angaza-referral.com
- Phone: +254 XXX XXX XXX
- Live Chat: Available 24/7
- Knowledge Base: Documentation Home
    """
    doc.add_paragraph(support_text)
    
    # Save the document
    doc.save('Angaza_Referral_System_Documentation.docx')
    print("\nDocumentation generation completed!")

if __name__ == '__main__':
    create_documentation() 